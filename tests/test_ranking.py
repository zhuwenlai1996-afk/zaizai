"""Validation tests for attendance ranking calculations."""
import os
import sys
import pytest

# Add project root to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from attendance_ranking.ranking import (
    read_parameters, read_roster, read_source_data,
    calculate_rankings, calculate_working_days, run
)
import openpyxl


INPUT_FILE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    '合肥基地氛围排名_严格排查修复_全名单个人排名.xlsx'
)


@pytest.fixture(scope='module')
def results():
    """Run the calculation once and share results across tests."""
    wb = openpyxl.load_workbook(INPUT_FILE, data_only=True)
    params = read_parameters(wb)
    roster = read_roster(wb)
    records = read_source_data(wb)
    return calculate_rankings(params, roster, records)


@pytest.fixture(scope='module')
def params():
    """Load parameters."""
    wb = openpyxl.load_workbook(INPUT_FILE, data_only=True)
    return read_parameters(wb)


class TestParameters:
    def test_working_day_count(self, params, results):
        """Working day count should be 14."""
        assert results['workday_count'] == 14

    def test_roster_count(self, results):
        """Total roster should have 121 people."""
        assert len(results['individual']) == 121


class TestIndividualRanking:
    def test_top_individual_name(self, results):
        """Top ranked individual should be 金永超."""
        top = results['individual'][0]
        assert top['name'] == '金永超'

    def test_top_individual_department(self, results):
        """Top ranked individual should be from 研发中心."""
        top = results['individual'][0]
        assert top['department'] == '研发中心'

    def test_top_individual_average(self, results):
        """Top ranked individual average should be ~14.53."""
        top = results['individual'][0]
        assert abs(top['average'] - 14.53) < 0.01

    def test_top_individual_workday_total(self, results):
        """Top ranked individual workday total should be 179.92."""
        top = results['individual'][0]
        assert abs(top['workday_total'] - 179.92) < 0.01

    def test_top_individual_valid_days(self, results):
        """Top ranked individual valid days should be 14."""
        top = results['individual'][0]
        assert top['valid_days'] == 14

    def test_top_individual_saturday_total(self, results):
        """Top ranked individual saturday total should be 23.5."""
        top = results['individual'][0]
        assert abs(top['saturday_total'] - 23.5) < 0.01


class TestDepartmentRanking:
    def test_top_zhineng_department(self, results):
        """Top 职能部门 should be 人事+行政+财务."""
        top = results['zhineng'][0]
        assert top['name'] == '人事+行政+财务'

    def test_top_zhineng_average(self, results):
        """Top 职能部门 average should be ~13.34."""
        top = results['zhineng'][0]
        assert abs(top['average'] - 13.3376) < 0.01

    def test_top_xiaoshou_department(self, results):
        """Top 销售部门 should be 整单技术."""
        top = results['xiaoshou'][0]
        assert top['name'] == '整单技术'

    def test_top_xiaoshou_average(self, results):
        """Top 销售部门 average should be ~13.33."""
        top = results['xiaoshou'][0]
        assert abs(top['average'] - 13.3349) < 0.01


class TestOutputGeneration:
    def test_output_file_creation(self, params, results, tmp_path):
        """Output Excel should be generated with correct sheets."""
        from attendance_ranking.ranking import generate_output
        output_path = str(tmp_path / 'test_ranking_output.xlsx')
        generate_output(params, results, output_path)
        assert os.path.exists(output_path)

        wb = openpyxl.load_workbook(output_path)
        assert '个人排名' in wb.sheetnames
        assert '部门排名' in wb.sheetnames
        assert '月度排名' in wb.sheetnames
        assert '检查项' in wb.sheetnames

    def test_output_individual_count(self, params, results, tmp_path):
        """Output 个人排名 should have 121 data rows."""
        from attendance_ranking.ranking import generate_output
        output_path = str(tmp_path / 'test_ranking_output.xlsx')
        generate_output(params, results, output_path)
        wb = openpyxl.load_workbook(output_path)
        ws = wb['个人排名']
        row_count = 0
        for row in ws.iter_rows(min_row=2, max_row=200, min_col=1, max_col=1, values_only=True):
            if row[0] is not None:
                row_count += 1
        assert row_count == 121
