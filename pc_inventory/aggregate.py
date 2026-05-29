"""汇总器: 把共享目录下所有采集 JSON 合并成一份按部门分组的 Excel.

用法::

    python -m pc_inventory aggregate -i \\\\server\\share\\pc_inventory\\data -o 部门电脑清单.xlsx

输出 Excel 包含:
    Sheet "总览"   - 全部记录, 按部门, 姓名 排序
    Sheet "<部门>" - 每个部门一个独立 Sheet (可用 --no-dept-sheets 关闭)
    Sheet "未分类" - 没填部门的记录
    Sheet "汇总"   - 按部门统计电脑数量
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from collections import defaultdict
from typing import Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# 常量
# ---------------------------------------------------------------------------
COLUMNS = [
    ("department", "部门", 16),
    ("real_name", "使用人", 10),
    ("username", "登录账号", 14),
    ("hostname", "主机名", 18),
    ("serial_number", "序列号", 18),
    ("primary_ip", "主 IP", 14),
    ("all_ips_str", "全部 IP", 24),
    ("mac", "MAC 地址", 18),
    ("manufacturer", "厂商", 12),
    ("model", "型号", 18),
    ("os", "操作系统", 28),
    ("collected_at", "采集时间", 19),
]

_INVALID_SHEET_CHARS = set(r'[]:*?/\\')


# ---------------------------------------------------------------------------
# 读取
# ---------------------------------------------------------------------------
def load_records(input_dir: str) -> List[dict]:
    if not os.path.isdir(input_dir):
        raise SystemExit(f"输入目录不存在: {input_dir}")

    records: List[dict] = []
    for name in sorted(os.listdir(input_dir)):
        if not name.lower().endswith(".json"):
            continue
        path = os.path.join(input_dir, name)
        try:
            with open(path, encoding="utf-8") as f:
                rec = json.load(f)
        except (OSError, json.JSONDecodeError) as exc:
            print(f"[警告] 跳过损坏文件 {name}: {exc}", file=sys.stderr)
            continue
        rec["_source_file"] = name
        # all_ips -> 字符串方便表格展示
        ips = rec.get("all_ips") or []
        rec["all_ips_str"] = ", ".join(ips) if isinstance(ips, list) else str(ips)
        records.append(rec)
    return records


def deduplicate_by_hostname(records: List[dict]) -> List[dict]:
    """同一 hostname 出现多次时只保留 collected_at 最新的那条."""
    by_host: Dict[str, dict] = {}
    for rec in records:
        host = (rec.get("hostname") or "").lower() or rec.get("_source_file", "")
        prev = by_host.get(host)
        if prev is None:
            by_host[host] = rec
            continue
        if (rec.get("collected_at") or "") > (prev.get("collected_at") or ""):
            by_host[host] = rec
    return list(by_host.values())


# ---------------------------------------------------------------------------
# 写出
# ---------------------------------------------------------------------------
_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_GRAY_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
_BORDER = Border(*(Side(style="thin") for _ in range(4)))


def _safe_sheet_name(name: str) -> str:
    cleaned = "".join("_" if c in _INVALID_SHEET_CHARS else c for c in name)
    return (cleaned or "未命名")[:31]  # Excel sheet name 上限 31


def _write_table(ws, records: List[dict]) -> None:
    # header
    for col_idx, (_, label, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, rec in enumerate(records, 2):
        for col_idx, (key, _, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=rec.get(key, ""))
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if row_idx % 2 == 0:
                cell.fill = _GRAY_FILL

    ws.freeze_panes = "A2"


def _write_summary(ws, by_dept: Dict[str, List[dict]]) -> None:
    headers = ["部门", "电脑数量"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    rows = sorted(
        ((dept or "未分类", len(recs)) for dept, recs in by_dept.items()),
        key=lambda x: (-x[1], x[0]),
    )
    for i, (dept, count) in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=dept).border = _BORDER
        ws.cell(row=i, column=2, value=count).border = _BORDER

    total = sum(c for _, c in rows)
    last = len(rows) + 2
    ws.cell(row=last, column=1, value="合计").font = Font(bold=True)
    ws.cell(row=last, column=2, value=total).font = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 12


def write_excel(records: List[dict], output_path: str, by_dept_sheets: bool = True) -> None:
    records = sorted(
        records,
        key=lambda r: (r.get("department") or "zzz_未分类", r.get("real_name") or "", r.get("hostname") or ""),
    )

    by_dept: Dict[str, List[dict]] = defaultdict(list)
    for rec in records:
        by_dept[rec.get("department") or ""].append(rec)

    wb = openpyxl.Workbook()

    # Sheet 1: 总览
    ws_all = wb.active
    ws_all.title = "总览"
    _write_table(ws_all, records)

    # Sheet: 每个部门
    if by_dept_sheets:
        used_names = {"总览"}
        # 排序: 有部门的按字母顺序, 未分类放最后
        for dept in sorted(by_dept.keys(), key=lambda d: (d == "", d)):
            label = dept or "未分类"
            sheet_name = _safe_sheet_name(label)
            base = sheet_name
            i = 2
            while sheet_name in used_names:
                sheet_name = f"{base}_{i}"[:31]
                i += 1
            used_names.add(sheet_name)
            ws = wb.create_sheet(sheet_name)
            _write_table(ws, by_dept[dept])

    # Sheet: 汇总
    ws_sum = wb.create_sheet("汇总")
    _write_summary(ws_sum, by_dept)

    wb.save(output_path)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        prog="pc_inventory aggregate",
        description="把共享目录里所有采集 JSON 汇总成一份 Excel.",
    )
    parser.add_argument("-i", "--input", required=True, help="共享数据目录, 内含 *.json")
    parser.add_argument("-o", "--output", default="部门电脑清单.xlsx", help="输出 Excel 路径")
    parser.add_argument(
        "--no-dept-sheets", action="store_true",
        help="不为每个部门生成单独 Sheet, 只保留总览 + 汇总.",
    )
    args = parser.parse_args(argv)

    print(f"[1/3] 加载 {args.input}")
    raw = load_records(args.input)
    print(f"      读到 {len(raw)} 条原始记录")

    print("[2/3] 按主机名去重 (保留最新)")
    records = deduplicate_by_hostname(raw)
    print(f"      去重后 {len(records)} 台电脑")

    print(f"[3/3] 写出 {args.output}")
    write_excel(records, args.output, by_dept_sheets=not args.no_dept_sheets)
    print("完成!")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
